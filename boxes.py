#!/usr/bin/env python3
"""Abstellraum inventory: one searchable page, read straight from the spreadsheet.

The sheet is the database. Edit boxes.xlsx, save, reload the page -- the mtime
change is noticed on the next request and the page is re-rendered.
"""

import html
import os
import threading
from datetime import datetime

import openpyxl
from flask import Flask, Response

# --- configuration ----------------------------------------------------------

HERE = os.path.dirname(os.path.abspath(__file__))
SHEET_PATH = os.environ.get("BOXES_SHEET", os.path.join(HERE, "boxes.xlsx"))
HOST = os.environ.get("BOXES_HOST", "127.0.0.1")
PORT = int(os.environ.get("BOXES_PORT", "8091"))

# Worksheet names inside the workbook.
SHEET_BOXES = "Inhalt"        # Regal | Box | ► | Inhalt, one item per row
SHEET_SHELVES = "Regale (2)"  # open shelves, a spatial grid with no ids

app = Flask(__name__)


# --- reading the spreadsheet ------------------------------------------------

def _text(value):
    """Cell value -> trimmed string. Integers must not render as '13.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _shelf_label(value):
    """Bare shelf values get the word Regal; RK and B stand on their own."""
    text = _text(value)
    if not text:
        return ""
    if text.replace(".", "", 1).isdigit():
        return "Regal " + _text(value)
    return text


def read_boxes(worksheet, formula_rows=frozenset()):
    """Walk the grouped layout: a non-empty Box cell starts a new box, and every
    row below it belongs to that box until the next one."""
    boxes = []
    current = None
    previous_number = 0

    for row in worksheet.iter_rows(min_row=3):
        shelf_cell, number_cell, _marker, content_cell = (list(row) + [None] * 4)[:4]
        number = _text(number_cell.value if number_cell else None)

        # A box number written as a formula (=B13+1) reads as blank when the
        # workbook was saved without cached results. Every such formula in this
        # sheet is 'previous + 1', so continue the sequence rather than silently
        # folding the box into the one above it.
        if not number and number_cell is not None and number_cell.row in formula_rows:
            number = str(previous_number + 1)

        if number:
            current = {
                "number": number,
                "shelf": _shelf_label(shelf_cell.value if shelf_cell else None),
                "items": [],
            }
            boxes.append(current)
            if number.isdigit():
                previous_number = int(number)

        item = _text(content_cell.value if content_cell else None)
        if current is not None and item:
            current["items"].append(item)

    return boxes


def read_shelves(worksheet):
    """The open-shelf grid: each column is one shelf unit, read top to bottom."""
    shelves = []
    for column in worksheet.iter_cols():
        contents = []
        for cell in column:
            item = _text(cell.value)
            # The same category repeated down a column means several levels of
            # it; for finding things once is enough.
            if item and item not in contents:
                contents.append(item)
        if contents:
            letter = column[0].column_letter
            shelves.append({"number": letter, "shelf": "offen", "items": contents})
    return shelves


def load_workbook_data(path):
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)

    # Rows whose box number is a formula, so read_boxes can tell 'empty' apart
    # from 'formula that was saved without a cached result'.
    formula_rows = {
        cell.row
        for row in formulas[SHEET_BOXES].iter_rows(min_row=3, min_col=2, max_col=2)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }

    boxes = read_boxes(values[SHEET_BOXES], formula_rows)
    shelves = read_shelves(values[SHEET_SHELVES]) if SHEET_SHELVES in values.sheetnames else []
    return boxes, shelves


# --- search haystack --------------------------------------------------------

_FOLD = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss", "é": "e", "è": "e", "à": "a"})


def haystack(entry, kind):
    """Everything a query may match, lowercased, plus umlaut- and hyphen-folded
    variants so 'muesli' finds Müsli and 'aufbewahrungsglaeser' finds the sheet's
    hyphenated spelling."""
    parts = [entry["number"], entry["shelf"], *entry["items"]]
    if kind == "box":
        parts += ["box " + entry["number"], "k" + entry["number"], "k " + entry["number"]]
    if kind == "shelf":
        parts.append("offenes regal")
    if not entry["items"]:
        parts.append("leer")

    plain = " ".join(parts).lower()
    folded = plain.translate(_FOLD)
    # Umlaut folding and hyphen removal have to combine, or the sheet's
    # "Aufbewahrungs-gläser" stays unreachable from "aufbewahrungsglaeser".
    variants = [plain, folded, plain.replace("-", ""), folded.replace("-", "")]

    unique = []
    for variant in variants:
        if variant not in unique:
            unique.append(variant)
    return " ".join(unique)


# --- rendering --------------------------------------------------------------

CSS = """
:root {
  color-scheme: light dark;
  --bg: #fdfdfc;
  --fg: #16181c;
  --muted: #666c76;
  --rule: #dfe1e5;
  --accent: #0b4ea2;
  --mark-bg: #cbe0f8;
  --mark-fg: #062c5c;
  --field: #ffffff;
}
@media (prefers-color-scheme: dark) {
  :root {
    --bg: #14161a;
    --fg: #eceef2;
    --muted: #9aa1ad;
    --rule: #2b2f36;
    --accent: #7db2f0;
    --mark-bg: #1d3f66;
    --mark-fg: #e4eefb;
    --field: #1c1f25;
  }
}
* { box-sizing: border-box; }
html { -webkit-text-size-adjust: 100%; }
body {
  margin: 0;
  background: var(--bg);
  color: var(--fg);
  font: 400 17px/1.45 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
  padding-bottom: 5rem;
}
.wrap { max-width: 40rem; margin: 0 auto; padding: 0 1.15rem; }

header { padding: 1.6rem 0 1rem; }
h1 { margin: 0; font-size: 1.45rem; font-weight: 650; letter-spacing: -0.01em; }
.stand { margin: 0.2rem 0 0; color: var(--muted); font-size: 0.85rem; }

.searchbar {
  position: sticky; top: 0; z-index: 5;
  background: var(--bg);
  padding: 0.5rem 0 0.7rem;
  border-bottom: 1px solid var(--rule);
}
input[type=search] {
  width: 100%;
  font: inherit;
  font-size: 17px;
  padding: 0.85rem 0.9rem;
  color: var(--fg);
  background: var(--field);
  border: 1px solid var(--rule);
  border-radius: 7px;
  -webkit-appearance: none;
  appearance: none;
}
input[type=search]::-webkit-search-decoration { -webkit-appearance: none; }
input[type=search]:focus-visible,
.sum:focus-visible {
  outline: 3px solid var(--accent);
  outline-offset: 2px;
}
.count { margin: 0.55rem 0 0; color: var(--muted); font-size: 0.85rem; }

h2 {
  margin: 2.4rem 0 0;
  padding-bottom: 0.4rem;
  font-size: 0.95rem;
  font-weight: 600;
  color: var(--muted);
  border-bottom: 1px solid var(--rule);
}

/* Box | Regal | Inhalt, plus a hanging count. The header row and every
   box row share these tracks, so the columns line up down the page. */
:root {
  --col-box: 2.8rem;
  --col-regal: 4.4rem;
  --col-gap: 0.7rem;
}
.headrow, .sum {
  display: grid;
  grid-template-columns: var(--col-box) var(--col-regal) 1fr auto;
  gap: 0 var(--col-gap);
  width: 100%;
  text-align: left;
}
.headrow {
  padding: 1.1rem 0 0.45rem;
  font-size: 0.78rem;
  font-weight: 600;
  color: var(--muted);
  border-bottom: 1px solid var(--rule);
}
.box { border-bottom: 1px solid var(--rule); }
.sum {
  align-items: baseline;
  min-height: 64px;
  padding: 0.85rem 0;
  font: inherit;
  color: inherit;
  background: none;
  border: 0;
  cursor: pointer;
}
.num {
  font-size: 1.6rem;
  font-weight: 600;
  font-variant-numeric: tabular-nums;
  letter-spacing: -0.02em;
}
.shelf {
  font-size: 0.83rem;
  font-weight: 600;
  color: var(--accent);
}
.preview.leer { color: var(--muted); font-style: italic; }
.more {
  color: var(--muted);
  font-size: 0.85rem;
  font-variant-numeric: tabular-nums;
}
/* Opened items hang under the Inhalt column, not under the box number. */
.items {
  display: none;
  margin: 0;
  padding: 0 0 1rem calc(var(--col-box) + var(--col-regal) + var(--col-gap) * 2);
  list-style: none;
}
.box.open .items { display: block; }
.items li { padding: 0.38rem 0; }
.items li::before { content: "– "; color: var(--muted); }
.box.open .num, .box.open .preview { color: var(--accent); }
mark { background: var(--mark-bg); color: var(--mark-fg); border-radius: 2px; padding: 0 0.1em; }
.none { padding: 2.5rem 0; color: var(--muted); }
[hidden] { display: none !important; }
"""

JS = """
(function () {
  var field = document.getElementById('q');
  var count = document.getElementById('count');
  var boxes = Array.prototype.slice.call(document.querySelectorAll('.box'));
  var total = boxes.length;
  var empty = document.getElementById('none');

  boxes.forEach(function (box) {
    var head = box.querySelector('.sum');
    if (!head) return;
    head.addEventListener('click', function () {
      var wasOpen = box.classList.contains('open');
      boxes.forEach(close);
      if (!wasOpen) open(box);
    });
  });

  function open(box) {
    box.classList.add('open');
    var head = box.querySelector('.sum');
    if (head) head.setAttribute('aria-expanded', 'true');
  }
  function close(box) {
    box.classList.remove('open');
    var head = box.querySelector('.sum');
    if (head) head.setAttribute('aria-expanded', 'false');
  }

  function fold(s) {
    return s.replace(/ä/g, 'ae').replace(/ö/g, 'oe').replace(/ü/g, 'ue').replace(/ß/g, 'ss');
  }

  // Rebuild a node's text, wrapping occurrences of the query in <mark>.
  // Built from text nodes, never innerHTML, so sheet content stays inert.
  function highlight(node, needle) {
    var text = node.dataset.text;
    node.textContent = '';
    if (!needle) { node.textContent = text; return; }
    var hay = text.toLowerCase(), at = 0, found;
    while ((found = hay.indexOf(needle, at)) !== -1) {
      node.appendChild(document.createTextNode(text.slice(at, found)));
      var m = document.createElement('mark');
      m.textContent = text.slice(found, found + needle.length);
      node.appendChild(m);
      at = found + needle.length;
    }
    node.appendChild(document.createTextNode(text.slice(at)));
  }

  function filter() {
    var raw = field.value.trim().toLowerCase();
    var variants = raw ? [raw, fold(raw), raw.replace(/-/g, ''), fold(raw).replace(/-/g, '')] : [];
    var hits = 0;

    boxes.forEach(function (box) {
      var hay = box.dataset.hay;
      var match = !raw || variants.some(function (v) { return hay.indexOf(v) !== -1; });
      box.hidden = !match;
      if (!match) {
        close(box);
        // Drop highlights from the previous query, or this row would still
        // carry them the next time it matches and is shown again.
        if (box.querySelector('mark')) {
          box.querySelectorAll('[data-text]').forEach(function (n) { highlight(n, ''); });
        }
        return;
      }
      hits++;
      // A match may be on an item that the collapsed row does not show, so
      // open it and point at the words that matched.
      if (raw) { open(box); } else { close(box); }
      box.querySelectorAll('[data-text]').forEach(function (n) { highlight(n, raw); });
    });

    document.querySelectorAll('section').forEach(function (s) {
      s.hidden = !s.querySelector('.box:not([hidden])');
    });
    empty.hidden = hits > 0;
    count.textContent = raw
      ? (hits === 1 ? '1 Treffer' : hits + ' Treffer')
      : total + ' Einträge';
  }

  field.addEventListener('input', filter);
  filter();
})();
"""


def esc(value):
    return html.escape(str(value), quote=True)


HEADROW = ('<div class="headrow" aria-hidden="true">'
           "<span>Box</span><span>Regal</span><span>Inhalt</span><span></span>"
           "</div>")


def render_entry(entry, kind):
    number = esc(entry["number"])
    shelf = esc(entry["shelf"]) if entry["shelf"] else "ohne Regal"
    items = entry["items"]
    hay = esc(haystack(entry, kind))

    out = ['<div class="box" data-hay="%s">' % hay]

    if not items:
        out.append('<div class="sum">')
        out.append('<span class="num">%s</span>' % number)
        out.append('<span class="shelf">%s</span>' % shelf)
        out.append('<span class="preview leer">leer</span>')
        out.append("</div>")
        out.append("</div>")
        return "".join(out)

    first = esc(items[0])
    out.append('<button class="sum" type="button" aria-expanded="false">')
    out.append('<span class="num">%s</span>' % number)
    out.append('<span class="shelf">%s</span>' % shelf)
    out.append('<span class="preview" data-text="%s">%s</span>' % (first, first))
    if len(items) > 1:
        out.append('<span class="more">+%d</span>' % (len(items) - 1))
    out.append("</button>")

    out.append('<ul class="items">')
    for item in items:
        text = esc(item)
        out.append('<li data-text="%s">%s</li>' % (text, text))
    out.append("</ul></div>")
    return "".join(out)


def render_page(boxes, shelves, stand):
    total = len(boxes) + len(shelves)
    parts = [
        "<!doctype html><html lang=de><head><meta charset=utf-8>",
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        "<title>Abstellraum</title>",
        '<meta name="color-scheme" content="light dark">',
        "<style>%s</style></head><body><div class=wrap>" % CSS,
        "<header><h1>Abstellraum</h1>",
        '<p class="stand">Stand %s</p></header>' % esc(stand),
        '<div class="searchbar">',
        '<input id="q" type="search" placeholder="Suchen…" aria-label="Inhalt suchen"',
        ' autocomplete="off" autocapitalize="off" autocorrect="off" spellcheck="false"',
        ' enterkeyhint="search">',
        '<p class="count" id="count" role="status">%d Einträge</p>' % total,
        "</div>",
        "<section>",
        HEADROW,
    ]
    parts += [render_entry(box, "box") for box in boxes]
    parts.append("</section>")

    if shelves:
        parts.append("<section><h2>Offene Regale</h2>" + HEADROW)
        parts += [render_entry(shelf, "shelf") for shelf in shelves]
        parts.append("</section>")

    parts.append('<p class="none" id="none" hidden>Nichts gefunden.</p>')
    parts.append("</div><script>%s</script></body></html>" % JS)
    return "".join(parts)


# --- cache ------------------------------------------------------------------

_lock = threading.Lock()
_cache = {"stamp": None, "html": None}


def current_page():
    """Render only when the spreadsheet has actually changed on disk."""
    info = os.stat(SHEET_PATH)
    stamp = (info.st_mtime_ns, info.st_size)

    with _lock:
        if _cache["stamp"] == stamp and _cache["html"] is not None:
            return _cache["html"]

        boxes, shelves = load_workbook_data(SHEET_PATH)
        stand = datetime.fromtimestamp(info.st_mtime).strftime("%d.%m.%Y %H:%M")
        page = render_page(boxes, shelves, stand)
        _cache["stamp"] = stamp
        _cache["html"] = page
        return page


@app.route("/")
def index():
    response = Response(current_page(), content_type="text/html; charset=utf-8")
    response.headers["Cache-Control"] = "no-store"
    return response


if __name__ == "__main__":
    from waitress import serve

    print("Abstellraum: %s -> http://%s:%d" % (SHEET_PATH, HOST, PORT), flush=True)
    serve(app, host=HOST, port=PORT, threads=4)
